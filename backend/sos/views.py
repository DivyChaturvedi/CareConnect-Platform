from rest_framework import generics, status
from rest_framework.response import Response
from rest_framework.permissions import IsAuthenticated
from .models import EmergencyCategory, SOSAlert
from .serializers import (
    EmergencyCategorySerializer,
    SOSAlertCreateSerializer,
    SOSAlertDetailSerializer,
    SOSAlertUpdateSerializer
)

from .utils import reverse_geocode


from notifications.services import notify_user
from users.models import User

from .models import EscalationLog

from rest_framework.generics import ListAPIView


class EmergencyCategoryListView(generics.ListAPIView):
    """Master data — sab active categories, resident ke SOS button screen ke liye"""
    queryset = EmergencyCategory.objects.filter(is_active=True)
    serializer_class = EmergencyCategorySerializer
    permission_classes = [IsAuthenticated]




from .services import route_sos_alert

from .models import EscalationConfig

from .tasks import escalate_to_secondary_guardian
from .services import route_sos_alert,broadcast_to_community

class SOSAlertCreateView(generics.CreateAPIView):
    serializer_class = SOSAlertCreateSerializer
    permission_classes = [IsAuthenticated]

    def create(self, request, *args, **kwargs):
        if request.user.role != "RESIDENT":
            return Response({"error": "Only residents can trigger SOS"}, status=status.HTTP_403_FORBIDDEN)

        serializer = self.get_serializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        alert = serializer.save()

        if alert.latitude and alert.longitude:
            address = reverse_geocode(alert.latitude, alert.longitude)
            if address:
                alert.address = address
                alert.save()

        # Stage 1: Primary Guardian ko turant notify karo
        route_sos_alert(alert)


        # Community broadcast — turant, parallel
        broadcast_to_community(alert)

        # Stage 2 schedule karo (Secondary Guardian, agar Stage 1 se response na aaye)
        config = EscalationConfig.objects.filter(is_active=True).first()
        window_minutes = config.response_window_minutes if config else 5
        escalate_to_secondary_guardian.apply_async(args=[alert.id], countdown=window_minutes * 60)

        return Response(
    {
        "success": True,
        "message": "SOS alert created successfully",
        "data": SOSAlertDetailSerializer(alert).data
    },
    status=status.HTTP_201_CREATED
)

    
class MySOSAlertsView(generics.ListAPIView):
    """Resident apne khud ke SOS alerts dekh sake (history)"""
    serializer_class = SOSAlertDetailSerializer
    permission_classes = [IsAuthenticated]

    def get_queryset(self):
        return SOSAlert.objects.filter(resident=self.request.user).order_by('-created_at')
    

class SOSAlertUpdateView(generics.UpdateAPIView):
    """SOS ke baad message add karna ya status update karna (resident ya responder)"""
    queryset = SOSAlert.objects.all()
    serializer_class = SOSAlertUpdateSerializer
    permission_classes = [IsAuthenticated]

    def update(self, request, *args, **kwargs):
        instance = self.get_object()
        if (
        request.user == instance.resident
        and "status" in request.data
        ):
            return Response(
            {
                "error": "Residents cannot update SOS status"
             },
        status=status.HTTP_400_BAD_REQUEST
        )
        if instance.resident != request.user and request.user.role not in ["ADMIN", "SECURITY"]:
            return Response({"error": "Not authorized to update this alert"}, status=status.HTTP_403_FORBIDDEN)

        serializer = self.get_serializer(instance, data=request.data, partial=True)
        serializer.is_valid(raise_exception=True)
        serializer.save()
        return Response(SOSAlertDetailSerializer(instance).data, status=status.HTTP_200_OK)



from django.utils import timezone

from rest_framework import serializers

class EmptySerializer(serializers.Serializer):
    pass

from rest_framework.views import APIView
from django.utils import timezone

class RespondToSOSView(APIView):
    permission_classes = [IsAuthenticated]

    def post(self, request, pk):
        alert = SOSAlert.objects.get(pk=pk)

        if alert.responded_by is not None:
            return Response(
                {
                    "success": True,
                    "message": "Already responded"
                },
                status=status.HTTP_200_OK
            )

        alert.responded_by = request.user
        alert.responded_at = timezone.now()
        alert.status = "active"
        alert.save()

        return Response(
            {
                "success": True,
                "message": "Response recorded, escalation cancelled"
            },
            status=status.HTTP_200_OK
        )


class SOSAlertGuardianDetailView(generics.RetrieveAPIView):
    """Koi bhi authenticated user (guardian/security) ek specific SOS alert ki detail dekh sake"""
    queryset = SOSAlert.objects.all()
    serializer_class = SOSAlertDetailSerializer
    permission_classes = [IsAuthenticated]


from .models import EscalationConfig
from .serializers import EscalationConfigSerializer
from users.permissions import IsAdmin

class EscalationConfigView(generics.RetrieveUpdateAPIView):
    """Admin response-time window set kare (ek hi active config maintain karte hain)"""
    serializer_class = EscalationConfigSerializer
    permission_classes = [IsAuthenticated, IsAdmin]

    def get_object(self):
        config, _ = EscalationConfig.objects.get_or_create(id=1, defaults={'response_window_minutes': 5})
        return config    
    




from .models import EscalationLog
from .serializers import EscalationLogSerializer
from users.permissions import IsAdminOrSecurity

class EscalationLogListView(generics.ListAPIView):
    """Admin/Security escalation history dekh sake"""
    queryset = EscalationLog.objects.all().order_by('-created_at')
    serializer_class = EscalationLogSerializer
    permission_classes = [IsAuthenticated, IsAdminOrSecurity]    


import math
from rest_framework.views import APIView
from .models import Broadcast, BroadcastRecipient, IncidentVisibilityConfig, VolunteerAvailability, SecurityStaff
from .serializers import BroadcastSerializer, BroadcastRecipientSerializer
from notifications.services import notify_user

def calculate_haversine_distance(lat1, lon1, lat2, lon2):
    if lat1 is None or lon1 is None or lat2 is None or lon2 is None:
        return float('inf')
    try:
        lat1, lon1, lat2, lon2 = map(math.radians, [float(lat1), float(lon1), float(lat2), float(lon2)])
        dlat = lat2 - lat1
        dlon = lon2 - lon1
        a = math.sin(dlat/2)**2 + math.cos(lat1) * math.cos(lat2) * math.sin(dlon/2)**2
        c = 2 * math.asin(math.sqrt(a))
        return 6371 * c  # Earth radius in kilometers
    except Exception:
        return float('inf')


class BroadcastListCreateView(generics.ListCreateAPIView):
    serializer_class = BroadcastSerializer
    permission_classes = [IsAuthenticated]

    def get_queryset(self):
        return Broadcast.objects.all().order_by('-created_at')

    def create(self, request, *args, **kwargs):
        incident_id = request.data.get('incident_id')
        if not incident_id:
            return Response({"error": "incident_id is required"}, status=status.HTTP_400_BAD_REQUEST)

        try:
            incident = SOSAlert.objects.get(id=incident_id)
        except SOSAlert.DoesNotExist:
            return Response({"error": "Incident not found"}, status=status.HTTP_404_NOT_FOUND)

        # Get radius
        config = IncidentVisibilityConfig.objects.first()
        default_radius = config.default_radius if config else 2.0
        max_radius = config.max_radius if config else 5.0

        radius = request.data.get('radius')
        if radius is not None:
            try:
                radius = float(radius)
                if radius > max_radius:
                    radius = max_radius
            except ValueError:
                radius = default_radius
        else:
            radius = default_radius

        # Create Broadcast
        broadcast = Broadcast.objects.create(
            incident=incident,
            radius=radius,
            status='pending'
        )

        recipients_to_notify = []

        # 1. Online Volunteers
        volunteers = VolunteerAvailability.objects.filter(status='online')
        for vol in volunteers:
            dist = calculate_haversine_distance(incident.latitude, incident.longitude, vol.lat, vol.lng)
            if dist <= radius:
                recipients_to_notify.append((vol.user, 'volunteer'))

        # 2. Active Security Staff
        security = SecurityStaff.objects.filter(is_active=True)
        for sec in security:
            dist = calculate_haversine_distance(incident.latitude, incident.longitude, sec.lat, sec.lng)
            if dist <= radius:
                recipients_to_notify.append((sec.user, 'security'))

        # 3. Society Members (other residents in same society)
        try:
            resident_profile = incident.resident.resident_profile
            society = resident_profile.flat.block.society if (resident_profile and resident_profile.flat and resident_profile.flat.block) else None
            if society:
                other_residents = User.objects.filter(
                    role='RESIDENT',
                    resident_profile__flat__block__society=society
                ).exclude(id=incident.resident.id)
                for res in other_residents:
                    recipients_to_notify.append((res, 'member'))
        except Exception:
            pass

        # De-duplicate recipients by user id
        seen_user_ids = set()
        final_recipients = []
        for user, role in recipients_to_notify:
            if user.id not in seen_user_ids:
                seen_user_ids.add(user.id)
                final_recipients.append((user, role))

        category_name = incident.category.name if incident.category else "Emergency"
        title = f"🆘 Community Alert: {category_name}"
        message = f"{incident.resident.first_name} {incident.resident.last_name} needs help nearby."
        if incident.address:
            message += f" Location: {incident.address}"

        created_recipients = []
        for user, role in final_recipients:
            recipient = BroadcastRecipient.objects.create(
                broadcast=broadcast,
                user=user,
                role=role,
                status='notified'
            )
            created_recipients.append(recipient)
            try:
                notify_user(user, title, message, notification_type="sos_alert", related_sos_id=incident.id)
            except Exception:
                recipient.status = 'failed'
                recipient.save()

        broadcast.status = 'sent' if created_recipients else 'failed'
        broadcast.save()

        return Response(BroadcastSerializer(broadcast).data, status=status.HTTP_201_CREATED)


class BroadcastDetailView(generics.RetrieveAPIView):
    queryset = Broadcast.objects.all()
    serializer_class = BroadcastSerializer
    permission_classes = [IsAuthenticated]


class BroadcastRecipientListView(generics.ListAPIView):
    serializer_class = BroadcastRecipientSerializer
    permission_classes = [IsAuthenticated]

    def get_queryset(self):
        broadcast_id = self.kwargs.get('pk')
        return BroadcastRecipient.objects.filter(broadcast_id=broadcast_id)


class IncidentVisibilityView(APIView):
    permission_classes = [IsAuthenticated]

    def get(self, request, pk):
        try:
            incident = SOSAlert.objects.get(id=pk)
        except SOSAlert.DoesNotExist:
            return Response({"error": "Incident not found"}, status=status.HTTP_404_NOT_FOUND)

        radii = [1.0, 2.0, 5.0]
        results = {}

        volunteers = VolunteerAvailability.objects.filter(status='online')
        security = SecurityStaff.objects.filter(is_active=True)

        society_member_count = 0
        try:
            resident_profile = incident.resident.resident_profile
            society = resident_profile.flat.block.society if (resident_profile and resident_profile.flat and resident_profile.flat.block) else None
            if society:
                society_member_count = User.objects.filter(
                    role='RESIDENT',
                    resident_profile__flat__block__society=society
                ).exclude(id=incident.resident.id).count()
        except Exception:
            pass

        for r in radii:
            vol_in_range = []
            sec_in_range = []
            
            for vol in volunteers:
                dist = calculate_haversine_distance(incident.latitude, incident.longitude, vol.lat, vol.lng)
                if dist <= r:
                    vol_in_range.append({
                        "id": vol.user.id,
                        "name": f"{vol.user.first_name} {vol.user.last_name}".strip(),
                        "distance_km": round(dist, 2)
                    })

            for sec in security:
                dist = calculate_haversine_distance(incident.latitude, incident.longitude, sec.lat, sec.lng)
                if dist <= r:
                    sec_in_range.append({
                        "id": sec.user.id,
                        "post": sec.post,
                        "name": f"{sec.user.first_name} {sec.user.last_name}".strip(),
                        "distance_km": round(dist, 2)
                    })

            results[f"{int(r)}km"] = {
                "volunteers_count": len(vol_in_range),
                "volunteers": vol_in_range,
                "security_count": len(sec_in_range),
                "security": sec_in_range,
                "members_count": society_member_count,
            }

        return Response({
            "incident_id": incident.id,
            "latitude": incident.latitude,
            "longitude": incident.longitude,
            "simulations": results
        }, status=status.HTTP_200_OK)    
    

from django.db.models import Avg, Count, F, ExpressionWrapper, DurationField
from django.utils import timezone
from datetime import timedelta
from notifications.models import NotificationDeliveryLog
from .serializers import SOSAlertMonitoringSerializer
from users.permissions import IsAdminOrSecurity


class AlertMonitoringListView(generics.ListAPIView):
    """Admin/Security saare alerts dekhein, filter ke saath"""
    serializer_class = SOSAlertMonitoringSerializer
    permission_classes = [IsAuthenticated, IsAdminOrSecurity]

    def get_queryset(self):
        queryset = SOSAlert.objects.select_related('resident', 'category').order_by('-created_at')
        status_filter = self.request.query_params.get('status')
        category_filter = self.request.query_params.get('category')

        if status_filter:
            queryset = queryset.filter(status=status_filter)
        if category_filter:
            queryset = queryset.filter(category_id=category_filter)

        return queryset



from rest_framework import serializers


class SosEmptySerializer(serializers.Serializer):
    pass
class DashboardStatsView(generics.GenericAPIView):
    """Charts ke liye combined stats"""
    serializer_class = SosEmptySerializer
    permission_classes = [IsAuthenticated, IsAdminOrSecurity]

    def get(self, request, *args, **kwargs):
        total_alerts = SOSAlert.objects.count()

        # Status breakdown
        status_breakdown = list(
            SOSAlert.objects.values('status').annotate(count=Count('id')).order_by('status')
        )

        # Delivery tracking (sent/delivered/failed per channel)
        delivery_breakdown = list(
            NotificationDeliveryLog.objects.values('channel', 'status').annotate(count=Count('id'))
        )

        # Response time (avg minutes between created_at and responded_at)
        responded_alerts = SOSAlert.objects.filter(responded_at__isnull=False).annotate(
            response_duration=ExpressionWrapper(
                F('responded_at') - F('created_at'), output_field=DurationField()
            )
        )
        avg_response = responded_alerts.aggregate(avg=Avg('response_duration'))['avg']
        avg_response_minutes = round(avg_response.total_seconds() / 60, 1) if avg_response else None

        # Alerts in last 7 days (for trend chart)
        seven_days_ago = timezone.now() - timedelta(days=7)
        recent_alerts = SOSAlert.objects.filter(created_at__gte=seven_days_ago)

        # Pre-fill all 7 days with 0 so the chart always has data points
        daily_counts = {}
        for i in range(7):
            day = (timezone.now() - timedelta(days=6 - i)).strftime('%Y-%m-%d')
            daily_counts[day] = 0
        for alert in recent_alerts:
            day_key = alert.created_at.strftime('%Y-%m-%d')
            daily_counts[day_key] = daily_counts.get(day_key, 0) + 1

        # Escalation rate
        escalated_count = SOSAlert.objects.filter(is_escalated=True).count()
        escalation_rate = round((escalated_count / total_alerts) * 100, 1) if total_alerts else 0




        # Response Time Distribution (buckets)
        buckets = {'0-2 min': 0, '2-5 min': 0, '5-10 min': 0, '10-20 min': 0, '20+ min': 0}
        for alert in responded_alerts:
            mins = alert.response_duration.total_seconds() / 60
            if mins <= 2:
                buckets['0-2 min'] += 1
            elif mins <= 5:
                buckets['2-5 min'] += 1
            elif mins <= 10:
                buckets['5-10 min'] += 1
            elif mins <= 20:
                buckets['10-20 min'] += 1
            else:
                buckets['20+ min'] += 1
        response_time_distribution = [{'range': k, 'count': v} for k, v in buckets.items()]

        # Escalation Funnel
        primary_notified = total_alerts
        no_response_count = SOSAlert.objects.filter(is_escalated=True).count()
        escalated_count_funnel = EscalationLog.objects.values('alert').distinct().count()
        resolved_after_escalation = SOSAlert.objects.filter(is_escalated=True, status__in=['resolved', 'closed']).count()

        escalation_funnel = [
            {'stage': 'Primary Notified', 'count': primary_notified},
            {'stage': 'No Response', 'count': no_response_count},
            {'stage': 'Escalated', 'count': escalated_count_funnel},
            {'stage': 'Resolved After Escalation', 'count': resolved_after_escalation},
        ]

        # Per-Channel Delivery Table
        channel_stats = []
        for channel_code, channel_label in NotificationDeliveryLog.CHANNEL_CHOICES:
            channel_logs = NotificationDeliveryLog.objects.filter(channel=channel_code)
            sent = channel_logs.filter(status='sent').count()
            failed = channel_logs.filter(status='failed').count()
            total = sent + failed
            delivery_rate = round((sent / total) * 100, 1) if total else 0
            channel_stats.append({
                'channel': channel_label,
                'sent': sent,
                'failed': failed,
                'delivery_rate': delivery_rate,
            })

        return Response({
            'total_alerts': total_alerts,
            'status_breakdown': status_breakdown,
            'delivery_breakdown': delivery_breakdown,
            'avg_response_minutes': avg_response_minutes,
            'escalation_rate': escalation_rate,
            'daily_trend': [{'date': k, 'count': v} for k, v in sorted(daily_counts.items())],
            'response_time_distribution': response_time_distribution,
            'escalation_funnel': escalation_funnel,
            'channel_stats': channel_stats,
            })

from .models import ResponseUpdate
class AcceptIncidentView(generics.UpdateAPIView):
    """Volunteer/Security ek incident accept kar sakta hai, duplicate-assignment prevent karte hue"""
    queryset = SOSAlert.objects.all()
    serializer_class = SosEmptySerializer
    permission_classes = [IsAuthenticated]

    def update(self, request, *args, **kwargs):
        if request.user.role not in ["VOLUNTEER", "SECURITY"]:
            return Response({"error": "Only volunteers or security can accept incidents"}, status=status.HTTP_403_FORBIDDEN)

        alert = self.get_object()

        if alert.assigned_to is not None:
            if alert.assigned_to_id == request.user.id:
                return Response({"message": "You have already accepted this incident"}, status=status.HTTP_200_OK)
            assigned_name = f"{alert.assigned_to.first_name} {alert.assigned_to.last_name}".strip()
            return Response(
                {"error": f"Already accepted by {assigned_name}"},
                status=status.HTTP_409_CONFLICT
            )

        alert.assigned_to = request.user
        alert.assigned_at = timezone.now()
        if alert.status in ['open', 'escalated']:
            alert.status = 'active'
        alert.save()

        ResponseUpdate.objects.create(
            alert=alert, actor=request.user, update_type='accepted',
            description=f"{request.user.first_name} {request.user.last_name} accepted the incident",
        )


        # Resident ko notify karo ki koi responder assign hua
        notify_user(
            alert.resident,
            "✅ Responder Assigned",
            f"{request.user.first_name} {request.user.last_name} is on the way to help you.",
            notification_type="approval",
            related_sos_id=alert.id,
            
        )
        

        return Response(SOSAlertDetailSerializer(alert).data, status=status.HTTP_200_OK)
    



class SOSAlertRetrieveView(generics.RetrieveAPIView):
    """Koi bhi authenticated user (Volunteer/Security/Guardian) ek specific alert dekh sake"""
    queryset = SOSAlert.objects.all()
    serializer_class = SOSAlertDetailSerializer
    permission_classes = [IsAuthenticated]





from django.utils import timezone
from .services import validate_status_transition
from .serializers import SOSStatusUpdateSerializer

class UpdateIncidentStatusView(generics.UpdateAPIView):
    """Status update karo lifecycle rules ke saath, closure notes ke saath"""
    queryset = SOSAlert.objects.all()
    serializer_class = SOSStatusUpdateSerializer
    permission_classes = [IsAuthenticated]

    def update(self, request, *args, **kwargs):
        alert = self.get_object()

        # Sirf resident (khud), assigned responder, ya admin/security update kar sake
        allowed_users = [alert.resident_id, alert.assigned_to_id]
        if request.user.id not in allowed_users and request.user.role not in ["ADMIN", "SECURITY"]:
            return Response({"error": "Not authorized to update this incident"}, status=status.HTTP_403_FORBIDDEN)

        data = request.data.copy() if hasattr(request.data, 'copy') else dict(request.data)
        if 'status' not in data and ('resolution_summary' in data or 'outcome' in data):
            data['status'] = 'closed'

        serializer = self.get_serializer(data=data)
        serializer.is_valid(raise_exception=True)
        new_status = serializer.validated_data['status']
        resolution_notes = serializer.validated_data.get('resolution_notes', '') or data.get('resolution_summary', '')

        outcome = serializer.validated_data.get('outcome', '') or data.get('outcome', '')
        if outcome:
            alert.outcome = outcome

        is_valid, error_msg = validate_status_transition(alert.status, new_status)
        if not is_valid:
            return Response({"error": error_msg}, status=status.HTTP_400_BAD_REQUEST)

        if new_status in ['resolved', 'closed'] and resolution_notes:
            alert.resolution_notes = resolution_notes
            if hasattr(alert, 'resolution_summary'):
                alert.resolution_summary = resolution_notes

        if new_status == 'closed':
            alert.closed_at = timezone.now()

        alert.status = new_status
        alert.save()

        ResponseUpdate.objects.create(
            alert=alert, actor=request.user, update_type='status_changed',
            description=f"Status changed to {new_status} by {request.user.first_name}",
        )

        return Response(SOSAlertDetailSerializer(alert).data, status=status.HTTP_200_OK)

    def post(self, request, *args, **kwargs):
        return self.update(request, *args, **kwargs)



from rest_framework import generics, status
from rest_framework.parsers import MultiPartParser, FormParser, JSONParser
from .models import ChatMessage
from .serializers import ChatMessageSerializer

class ChatHistoryView(generics.ListCreateAPIView):
    serializer_class = ChatMessageSerializer
    permission_classes = [IsAuthenticated]

    def get_queryset(self):
        # Swagger schema generation
        if getattr(self, "swagger_fake_view", False):
            return ChatMessage.objects.none()

        alert_id = self.kwargs.get("pk")

        if not alert_id:
            return ChatMessage.objects.none()

        return ChatMessage.objects.filter(alert_id=alert_id)

    def perform_create(self, serializer):
        alert_id = self.kwargs.get("pk")
        alert = SOSAlert.objects.get(id=alert_id)
        msg = serializer.save(alert=alert, sender=self.request.user)

        try:
            from channels.layers import get_channel_layer
            from asgiref.sync import async_to_sync
            channel_layer = get_channel_layer()
            if channel_layer:
                async_to_sync(channel_layer.group_send)(
                    f'chat_{alert_id}',
                    {
                        'type': 'chat_message',
                        'id': msg.id,
                        'message': msg.message,
                        'message_type': msg.message_type,
                        'image_url': msg.image_url,
                        'sender_id': self.request.user.id,
                        'sender_name': f"{self.request.user.first_name} {self.request.user.last_name}".strip(),
                        'sender_role': self.request.user.role,
                        'created_at': msg.created_at.isoformat(),
                    }
                )
        except Exception as e:
            print(f"[CHAT BROADCAST ERROR] {e}")


from users.models import GuardianProfile
from .models import ChatMessage, ChatReadStatus


class ChatParticipantsView(generics.GenericAPIView):
    permission_classes = [IsAuthenticated]
    serializer_class = SosEmptySerializer
    queryset = SOSAlert.objects.none()

    def get(self, request, *args, **kwargs):
        if getattr(self, "swagger_fake_view", False):
            return Response([])
        alert = SOSAlert.objects.get(id=kwargs['pk'])
        participants = []
        seen_ids = set()

        def add_user(user, role_label):
            if user and user.id not in seen_ids:
                seen_ids.add(user.id)
                participants.append({
                    'id': user.id,
                    'name': f"{user.first_name} {user.last_name}".strip(),
                    'role': role_label,
                })

        add_user(alert.resident, 'RESIDENT')
        if alert.assigned_to:
            add_user(alert.assigned_to, alert.assigned_to.role)

        guardians = GuardianProfile.objects.filter(
            linked_resident_phone=alert.resident.phone_number, is_verified=True
        ).select_related('user')
        for g in guardians:
            add_user(g.user, 'GUARDIAN')

        # Jinhone message bheja hai wo bhi participants hain
        senders = ChatMessage.objects.filter(alert=alert).values_list('sender', flat=True).distinct()
        from users.models import User
        for sender_id in senders:
            u = User.objects.filter(id=sender_id).first()
            if u:
                add_user(u, u.role)

        return Response(participants)


class ChatImageUploadView(generics.GenericAPIView):
    permission_classes = [IsAuthenticated]
    parser_classes = [MultiPartParser, FormParser]
    serializer_class = SosEmptySerializer
    queryset = SOSAlert.objects.none()

    def post(self, request, *args, **kwargs):
        image = request.FILES.get('image')
        if not image:
            return Response({"error": "No image provided"}, status=status.HTTP_400_BAD_REQUEST)

        import os
        from django.conf import settings as dj_settings
        from django.core.files.storage import default_storage

        filename = f"chat_images/{kwargs['pk']}_{image.name}"
        path = default_storage.save(filename, image)
        url = request.build_absolute_uri(dj_settings.MEDIA_URL + path)

        return Response({"url": url}, status=status.HTTP_201_CREATED)





from .models import ResponseUpdate
from .serializers import ResponseUpdateSerializer


class ResponseUpdatesListView(generics.ListAPIView):
    serializer_class = ResponseUpdateSerializer
    permission_classes = [IsAuthenticated]

    def get_queryset(self):
        if getattr(self, 'swagger_fake_view', False):
            return ResponseUpdate.objects.none()

        return ResponseUpdate.objects.filter(
            alert_id=self.kwargs.get('pk')
        )


from .models import IncidentAttachment
from .serializers import IncidentAttachmentSerializer, IncidentClosureSerializer
from .services import VALID_TRANSITIONS, validate_status_transition
from rest_framework.parsers import MultiPartParser, FormParser, JSONParser


class IncidentTransitionsView(APIView):
    """GET /api/incidents/{id}/transitions/ - Get allowed next states"""
    permission_classes = [IsAuthenticated]

    def get(self, request, pk):
        try:
            alert = SOSAlert.objects.get(pk=pk)
        except SOSAlert.DoesNotExist:
            return Response({"error": "Incident not found"}, status=status.HTTP_404_NOT_FOUND)

        valid_next = VALID_TRANSITIONS.get(alert.status, [])
        return Response({
            "incident_id": alert.id,
            "current_status": alert.status,
            "allowed_transitions": valid_next,
        }, status=status.HTTP_200_OK)


class IncidentClosureView(APIView):
    """POST /api/incidents/{id}/closure/ - Submit resolution notes & closure details"""
    permission_classes = [IsAuthenticated]

    def post(self, request, pk):
        try:
            alert = SOSAlert.objects.get(pk=pk)
        except SOSAlert.DoesNotExist:
            return Response({"error": "Incident not found"}, status=status.HTTP_404_NOT_FOUND)

        # Authorization: Resident, Assigned responder, Admin/Security
        allowed_users = [alert.resident_id, alert.assigned_to_id]
        if request.user.id not in allowed_users and request.user.role not in ["ADMIN", "SECURITY"]:
            return Response({"error": "Not authorized to submit closure details for this incident"}, status=status.HTTP_403_FORBIDDEN)

        serializer = IncidentClosureSerializer(data=request.data)
        serializer.is_valid(raise_exception=True)

        resolution_summary = serializer.validated_data['resolution_summary']
        outcome = serializer.validated_data['outcome']
        resolution_notes = serializer.validated_data.get('resolution_notes', resolution_summary)
        res_datetime = serializer.validated_data.get('resolution_datetime', timezone.now())

        # Validate status transition if not already closed
        if alert.status != 'closed':
            is_valid, error_msg = validate_status_transition(alert.status, 'closed')
            if not is_valid:
                return Response({"error": error_msg}, status=status.HTTP_400_BAD_REQUEST)

        alert.resolution_summary = resolution_summary
        alert.resolution_notes = resolution_notes
        alert.outcome = outcome
        alert.closed_at = res_datetime or timezone.now()
        alert.status = 'closed'
        alert.save()

        ResponseUpdate.objects.create(
            alert=alert,
            actor=request.user,
            update_type='status_changed',
            description=f"Incident CLOSED with outcome '{outcome}': {resolution_summary[:100]}",
        )

        return Response({
            "success": True,
            "message": "Closure details submitted and incident closed successfully",
            "data": SOSAlertDetailSerializer(alert).data
        }, status=status.HTTP_200_OK)


class IncidentAttachmentsView(APIView):
    """POST/GET /api/incidents/{id}/attachments/ - Upload / List evidence & attachments"""
    permission_classes = [IsAuthenticated]
    parser_classes = [MultiPartParser, FormParser, JSONParser]

    def get(self, request, pk):
        try:
            alert = SOSAlert.objects.get(pk=pk)
        except SOSAlert.DoesNotExist:
            return Response({"error": "Incident not found"}, status=status.HTTP_404_NOT_FOUND)

        attachments = IncidentAttachment.objects.filter(incident=alert)
        serializer = IncidentAttachmentSerializer(attachments, many=True)
        return Response(serializer.data, status=status.HTTP_200_OK)

    def post(self, request, pk):
        try:
            alert = SOSAlert.objects.get(pk=pk)
        except SOSAlert.DoesNotExist:
            return Response({"error": "Incident not found"}, status=status.HTTP_404_NOT_FOUND)

        uploaded_files = request.FILES.getlist('file') or request.FILES.getlist('attachments')
        if not uploaded_files:
            file_obj = request.FILES.get('file')
            if file_obj:
                uploaded_files = [file_obj]

        if not uploaded_files:
            return Response({"error": "No file uploaded"}, status=status.HTTP_400_BAD_REQUEST)

        created_attachments = []
        for file_obj in uploaded_files:
            attachment = IncidentAttachment.objects.create(
                incident=alert,
                file=file_obj,
                file_name=file_obj.name,
                file_type=getattr(file_obj, 'content_type', 'application/octet-stream'),
                file_size=getattr(file_obj, 'size', 0),
                uploaded_by=request.user
            )
            created_attachments.append(attachment)

        ResponseUpdate.objects.create(
            alert=alert,
            actor=request.user,
            update_type='note',
            description=f"Uploaded {len(created_attachments)} attachment(s)",
        )

        return Response({
            "success": True,
            "message": f"Successfully uploaded {len(created_attachments)} file(s)",
            "attachments": IncidentAttachmentSerializer(created_attachments, many=True).data
        }, status=status.HTTP_201_CREATED)


class IncidentTimelineView(APIView):
    """GET /api/incidents/{id}/timeline/ - Get lifecycle timeline and audit trail"""
    permission_classes = [IsAuthenticated]

    def get(self, request, pk):
        try:
            alert = SOSAlert.objects.get(pk=pk)
        except SOSAlert.DoesNotExist:
            return Response({"error": "Incident not found"}, status=status.HTTP_404_NOT_FOUND)

        events = []
        # 1. Created Event
        resident_name = f"{alert.resident.first_name} {alert.resident.last_name}".strip()
        category_name = alert.category.name if alert.category else "Emergency"
        events.append({
            "id": f"evt-created-{alert.id}",
            "type": "created",
            "title": "Incident Triggered",
            "status": "open",
            "actor": resident_name,
            "description": f"Incident {category_name} created by {resident_name}",
            "timestamp": alert.created_at,
        })

        # 2. Response Updates
        updates = ResponseUpdate.objects.filter(alert=alert).order_by('created_at')
        for update in updates:
            actor_name = f"{update.actor.first_name} {update.actor.last_name}".strip() if update.actor else "System"
            events.append({
                "id": f"evt-update-{update.id}",
                "type": update.update_type,
                "title": update.update_type.replace("_", " ").title(),
                "status": alert.status,
                "actor": actor_name,
                "description": update.description,
                "timestamp": update.created_at,
            })

        # 3. Escalation Logs
        escalations = EscalationLog.objects.filter(alert=alert).order_by('created_at')
        for esc in escalations:
            events.append({
                "id": f"evt-esc-{esc.id}",
                "type": "escalated",
                "title": "Incident Escalated",
                "status": "escalated",
                "actor": "Escalation Engine",
                "description": f"Escalated to {esc.escalated_to_name} ({esc.reason})",
                "timestamp": esc.created_at,
            })

        # Sort all timeline events chronologically
        events.sort(key=lambda x: x["timestamp"])

        return Response({
            "incident_id": alert.id,
            "current_status": alert.status,
            "timeline": events,
        }, status=status.HTTP_200_OK)


class IncidentStatusPostView(APIView):
    """POST /api/incidents/{id}/status/ - Update lifecycle status via POST"""
    permission_classes = [IsAuthenticated]

    def post(self, request, pk):
        try:
            alert = SOSAlert.objects.get(pk=pk)
        except SOSAlert.DoesNotExist:
            return Response({"error": "Incident not found"}, status=status.HTTP_404_NOT_FOUND)

        allowed_users = [alert.resident_id, alert.assigned_to_id]
        if request.user.id not in allowed_users and request.user.role not in ["ADMIN", "SECURITY"]:
            return Response({"error": "Not authorized to update this incident"}, status=status.HTTP_403_FORBIDDEN)

        new_status = request.data.get('status')
        resolution_notes = request.data.get('resolution_notes', '')

        if not new_status:
            return Response({"error": "status is required"}, status=status.HTTP_400_BAD_REQUEST)

        is_valid, error_msg = validate_status_transition(alert.status, new_status)
        if not is_valid:
            return Response({"error": error_msg}, status=status.HTTP_400_BAD_REQUEST)

        if new_status == 'resolved' and resolution_notes:
            alert.resolution_notes = resolution_notes

        if new_status == 'closed':
            alert.closed_at = timezone.now()
            if resolution_notes:
                alert.resolution_notes = resolution_notes

        alert.status = new_status
        alert.save()

        ResponseUpdate.objects.create(
            alert=alert, actor=request.user, update_type='status_changed',
            description=f"Status changed to {new_status} by {request.user.first_name}",
        )

        return Response(SOSAlertDetailSerializer(alert).data, status=status.HTTP_200_OK)


from .serializers import SecurityDashboardAlertSerializer

class SecurityDashboardView(generics.ListAPIView):
    """Security ke liye active incidents ki list"""
    serializer_class = SecurityDashboardAlertSerializer
    permission_classes = [IsAuthenticated]

    def get_queryset(self):
        if self.request.user.role not in ["SECURITY", "ADMIN"]:
            return SOSAlert.objects.none()
        return SOSAlert.objects.exclude(status__in=['closed']).select_related('resident', 'category', 'assigned_to').order_by('-created_at')


class SecurityReportSummaryView(generics.GenericAPIView):
    """Security ke liye quick summary stats"""
    permission_classes = [IsAuthenticated]
    serializer_class = SosEmptySerializer

    def get(self, request, *args, **kwargs):
        if request.user.role not in ["SECURITY", "ADMIN"]:
            return Response({"error": "Not authorized"}, status=status.HTTP_403_FORBIDDEN)

        total_active = SOSAlert.objects.exclude(status__in=['closed', 'resolved']).count()
        assigned_to_me = SOSAlert.objects.filter(assigned_to=request.user).exclude(status__in=['closed', 'resolved']).count()
        resolved_today = SOSAlert.objects.filter(
            status__in=['resolved', 'closed'],
            updated_at__date=timezone.now().date(),
        ).count()
        unassigned = SOSAlert.objects.filter(assigned_to__isnull=True).exclude(status__in=['closed', 'resolved']).count()

        return Response({
            'total_active': total_active,
            'assigned_to_me': assigned_to_me,
            'resolved_today': resolved_today,
            'unassigned': unassigned,
        })





from django.utils.dateparse import parse_date
from .services import get_alert_society


def apply_report_filters(request):
    """Query params se filtered queryset banata hai — reusable helper"""
    queryset = SOSAlert.objects.select_related('resident', 'category').all()

    date_from = request.query_params.get('date_from')
    date_to = request.query_params.get('date_to')
    society_id = request.query_params.get('society')
    category_id = request.query_params.get('category')

    if date_from:
        parsed = parse_date(date_from)
        if parsed:
            queryset = queryset.filter(created_at__date__gte=parsed)
    if date_to:
        parsed = parse_date(date_to)
        if parsed:
            queryset = queryset.filter(created_at__date__lte=parsed)
    if category_id:
        queryset = queryset.filter(category_id=category_id)
    if society_id:
        queryset = queryset.filter(resident__resident_profile__flat__block__society_id=society_id)

    return queryset


class ReportSummaryView(generics.GenericAPIView):
    """Filtered incident statistics + society-wise breakdown"""
    serializer_class = SosEmptySerializer
    permission_classes = [IsAuthenticated, IsAdminOrSecurity]

    def get(self, request, *args, **kwargs):
        queryset = apply_report_filters(request)

        total = queryset.count()
        status_breakdown = list(queryset.values('status').annotate(count=Count('id')))
        category_breakdown = list(
            queryset.values('category__name').annotate(count=Count('id')).order_by('-count')
        )

        responded = queryset.filter(responded_at__isnull=False).annotate(
            response_duration=ExpressionWrapper(F('responded_at') - F('created_at'), output_field=DurationField())
        )
        avg_response = responded.aggregate(avg=Avg('response_duration'))['avg']
        avg_response_minutes = round(avg_response.total_seconds() / 60, 1) if avg_response else None

        # Society-wise breakdown
        society_counts = {}
        for alert in queryset.select_related('resident__resident_profile__flat__block__society'):
            society = get_alert_society(alert)
            name = society.name if society else "Unmapped"
            society_counts[name] = society_counts.get(name, 0) + 1

        society_breakdown = [{'society': k, 'count': v} for k, v in sorted(society_counts.items(), key=lambda x: -x[1])]

        return Response({
            'total_alerts': total,
            'status_breakdown': status_breakdown,
            'category_breakdown': category_breakdown,
            'avg_response_minutes': avg_response_minutes,
            'society_breakdown': society_breakdown,
        })





from django.http import HttpResponse
import openpyxl
from openpyxl.styles import Font, PatternFill


class ExportExcelReportView(generics.GenericAPIView):
    serializer_class = SosEmptySerializer
    permission_classes = [IsAuthenticated, IsAdminOrSecurity]

    def get(self, request, *args, **kwargs):
        queryset = apply_report_filters(request).order_by('-created_at')

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Incident Report"

        headers = ['ID', 'Resident', 'Category', 'Status', 'Address', 'Created At', 'Responded At']
        ws.append(headers)
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(start_color="4F46E5", end_color="4F46E5", fill_type="solid")

        for alert in queryset:
            ws.append([
                alert.id,
                f"{alert.resident.first_name} {alert.resident.last_name}",
                alert.category.name if alert.category else "N/A",
                alert.status,
                alert.address or "N/A",
                alert.created_at.strftime('%Y-%m-%d %H:%M'),
                alert.responded_at.strftime('%Y-%m-%d %H:%M') if alert.responded_at else "N/A",
            ])

        for col in ws.columns:
            max_length = max(len(str(cell.value)) for cell in col)
            ws.column_dimensions[col[0].column_letter].width = max_length + 4

        response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        response['Content-Disposition'] = 'attachment; filename="incident_report.xlsx"'
        wb.save(response)
        return response




from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import cm
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet
import io


class ExportPdfReportView(generics.GenericAPIView):
    serializer_class = SosEmptySerializer
    permission_classes = [IsAuthenticated, IsAdminOrSecurity]

    def get(self, request, *args, **kwargs):
        queryset = apply_report_filters(request).order_by('-created_at')

        buffer = io.BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=A4)
        styles = getSampleStyleSheet()
        elements = []

        elements.append(Paragraph("CareConnect Incident Report", styles['Title']))
        elements.append(Spacer(1, 0.5 * cm))
        elements.append(Paragraph(f"Total Incidents: {queryset.count()}", styles['Normal']))
        elements.append(Spacer(1, 0.5 * cm))

        data = [['ID', 'Resident', 'Category', 'Status', 'Created At']]
        for alert in queryset[:200]:  # limit for performance
            data.append([
                str(alert.id),
                f"{alert.resident.first_name} {alert.resident.last_name}",
                alert.category.name if alert.category else "N/A",
                alert.status,
                alert.created_at.strftime('%Y-%m-%d %H:%M'),
            ])

        table = Table(data, repeatRows=1)
        table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#4F46E5')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('FONTSIZE', (0, 0), (-1, -1), 8),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F3F4F6')]),
        ]))
        elements.append(table)

        doc.build(elements)
        buffer.seek(0)

        response = HttpResponse(buffer, content_type='application/pdf')
        response['Content-Disposition'] = 'attachment; filename="incident_report.pdf"'
        return response





    




